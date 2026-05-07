"""Course Knowledge Base — Qdrant-backed per-course RAG.

Each course has its own dedicated Qdrant collection named:

    kb_{safe_institute_id}_{safe_course_id}

where "safe" means lower-cased with hyphens replaced by underscores.

This gives complete isolation between courses: no payload filtering is
required during search, and dropping a course's KB is a single collection
delete rather than a filtered bulk-delete.

Point payload schema:
    content      – raw chunk text (returned to the agent)
    page         – source page number (1-based)
    chunk_index  – position within the page
    content_id   – UUID of the ModuleContent record (used for upsert / delete)
    title        – content title
"""

import io
import logging
import time
import uuid
from pathlib import PurePosixPath
from threading import Lock
from typing import Optional
from urllib.parse import urlparse

import httpx
from fastembed import TextEmbedding

from app.config import QDRANT_API_KEY, QDRANT_URL

logger = logging.getLogger(__name__)

EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
EMBEDDING_DIM = 384
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 200
BATCH_SIZE = 50
MAX_CONTENT_CHARS = 500  # truncate chunks to limit tokens sent to model

# ── Search result cache ────────────────────────────────────────────────
_CACHE_TTL = 300  # seconds
_CACHE_MAX = 256
_search_cache: dict[tuple, tuple[float, list]] = {}
_search_cache_lock = Lock()


def _cache_get(key: tuple) -> list | None:
    with _search_cache_lock:
        entry = _search_cache.get(key)
    if entry and (time.monotonic() - entry[0]) < _CACHE_TTL:
        return entry[1]
    return None


def _cache_set(key: tuple, results: list) -> None:
    with _search_cache_lock:
        _search_cache[key] = (time.monotonic(), results)
        if len(_search_cache) > _CACHE_MAX:
            oldest = sorted(_search_cache.items(), key=lambda x: x[1][0])
            for k, _ in oldest[:64]:
                del _search_cache[k]

_embedding_model: Optional[TextEmbedding] = None


def _get_embedding_model() -> TextEmbedding:
    global _embedding_model
    if _embedding_model is None:
        logger.info(f"Loading FastEmbed model: {EMBEDDING_MODEL}")
        _embedding_model = TextEmbedding(EMBEDDING_MODEL)
    return _embedding_model


def _qdrant_headers() -> dict:
    headers = {"Content-Type": "application/json"}
    if QDRANT_API_KEY:
        headers["api-key"] = QDRANT_API_KEY
    return headers


def _qdrant_request(method: str, path: str, **kwargs) -> dict:
    url = f"{QDRANT_URL}{path}"
    with httpx.Client(timeout=30) as client:
        r = client.request(method, url, headers=_qdrant_headers(), **kwargs)
        r.raise_for_status()
        return r.json()


def collection_name(institute_id: str, course_id: str) -> str:
    """Return the Qdrant collection name for a given institute + course pair.

    Example: "kb_abc123_def456" (hyphens replaced with underscores).
    """
    safe_inst = institute_id.lower().replace("-", "_")
    safe_course = course_id.lower().replace("-", "_")
    return f"kb_{safe_inst}_{safe_course}"


def ensure_collection(col_name: str) -> None:
    """Create a Qdrant collection if it does not already exist."""
    result = _qdrant_request("GET", "/collections")
    existing = [c["name"] for c in result["result"]["collections"]]
    if col_name not in existing:
        _qdrant_request(
            "PUT",
            f"/collections/{col_name}",
            json={
                "vectors": {"size": EMBEDDING_DIM, "distance": "Cosine"},
                "hnsw_config": {"m": 16, "ef_construct": 100},
                "optimizers_config": {"default_segment_number": 2},
            },
        )
        logger.info(f"Created Qdrant collection: {col_name!r}")


# ── Text extraction ────────────────────────────────────────────────────


def _extract_text_pdf(data: bytes) -> list[dict]:
    """Return [{page, text}] from PDF bytes using PyMuPDF."""
    import fitz  # PyMuPDF

    doc = fitz.open(stream=data, filetype="pdf")
    pages = []
    for i, page in enumerate(doc):
        text = page.get_text().strip()
        if text:
            pages.append({"page": i + 1, "text": text})
    doc.close()
    return pages


def _extract_text_docx(data: bytes) -> list[dict]:
    """Return [{page, text}] from DOCX bytes using python-docx."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    full_text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [{"page": 1, "text": full_text}] if full_text else []


def _extract_text_pptx(data: bytes) -> list[dict]:
    """Return [{page, text}] from PPTX bytes — one entry per slide."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    pages = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = " ".join(run.text for run in para.runs if run.text.strip())
                    if line.strip():
                        lines.append(line.strip())
        if lines:
            pages.append({"page": slide_num, "text": "\n".join(lines)})
    return pages


def _extract_text_xlsx(data: bytes) -> list[dict]:
    """Return [{page, text}] from XLSX bytes — one entry per sheet."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    pages = []
    for sheet_num, sheet in enumerate(wb.worksheets, start=1):
        lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append("\t".join(cells))
        if lines:
            pages.append({"page": sheet_num, "text": "\n".join(lines)})
    wb.close()
    return pages


def _chunk_text(pages: list[dict]) -> list[dict]:
    """Split page text into overlapping fixed-size chunks."""
    chunks = []
    for page_info in pages:
        text = page_info["text"]
        page_num = page_info["page"]
        start = 0
        chunk_idx = 0
        while start < len(text):
            chunk = text[start : start + CHUNK_SIZE]
            chunks.append({"text": chunk, "page": page_num, "chunk_index": chunk_idx})
            start += CHUNK_SIZE - CHUNK_OVERLAP
            chunk_idx += 1
    return chunks


# ── Public API ─────────────────────────────────────────────────────────


def index_content(
    institute_id: str,
    course_id: str,
    course_name: str,
    content_id: str,
    file_url: str,
    file_type: str,
    title: str,
) -> int:
    """Download, extract, embed and upsert a document into the course KB collection.

    The collection is named ``kb_{institute_id}_{course_id}`` and is created
    automatically on first use.  Existing points for ``content_id`` are deleted
    first so re-uploads remain idempotent.

    Returns:
        Number of Qdrant points upserted, or 0 if nothing was indexed.
    """
    col = collection_name(institute_id, course_id)

    # 1. Download
    logger.info(f"Downloading content {content_id!r} from {file_url!r}")
    with httpx.Client(timeout=60, follow_redirects=True) as client:
        resp = client.get(file_url)
        resp.raise_for_status()
        file_bytes = resp.content

    # 2. Extract text — derive actual format from the URL file extension
    ft = file_type.lower()
    url_ext = PurePosixPath(urlparse(file_url).path).suffix.lstrip(".").lower()
    if ft == "pdf" or url_ext == "pdf":
        pages = _extract_text_pdf(file_bytes)
    elif url_ext in ("pptx", "ppt"):
        pages = _extract_text_pptx(file_bytes)
    elif url_ext in ("xlsx", "xls", "ods"):
        pages = _extract_text_xlsx(file_bytes)
    elif ft in ("document", "docx", "word") or url_ext in ("docx", "doc", "odt"):
        pages = _extract_text_docx(file_bytes)
    else:
        logger.warning(f"Unsupported file type for KB indexing: {file_type!r} (ext: {url_ext!r})")
        return 0

    if not pages:
        logger.warning(f"No text extracted from content {content_id!r}")
        return 0

    # 3. Chunk
    chunks = _chunk_text(pages)
    if not chunks:
        return 0

    # 4. Embed
    model = _get_embedding_model()
    embeddings = [e.tolist() for e in model.embed([c["text"] for c in chunks])]

    # 5. Delete existing points for this content (idempotent)
    try:
        _qdrant_request(
            "POST",
            f"/collections/{col}/points/delete",
            json={"filter": {"must": [{"key": "content_id", "match": {"value": content_id}}]}},
        )
    except Exception:
        pass  # Collection may not exist yet — created below

    # 6. Ensure collection exists, then upsert in batches
    ensure_collection(col)

    total = 0
    for i in range(0, len(chunks), BATCH_SIZE):
        batch_chunks = chunks[i : i + BATCH_SIZE]
        batch_embeddings = embeddings[i : i + BATCH_SIZE]
        points = [
            {
                "id": str(uuid.uuid5(uuid.NAMESPACE_DNS, f"{content_id}_chunk_{i + j}")),
                "vector": emb,
                "payload": {
                    "content": chunk["text"],
                    "page": chunk["page"],
                    "chunk_index": chunk["chunk_index"],
                    "content_id": content_id,
                    "title": title,
                },
            }
            for j, (chunk, emb) in enumerate(zip(batch_chunks, batch_embeddings))
        ]
        _qdrant_request("PUT", f"/collections/{col}/points", json={"points": points})
        total += len(points)

    logger.info(f"Indexed {total} points for content {content_id!r} into collection {col!r} (course: {course_name!r})")
    return total


def delete_content(institute_id: str, course_id: str, content_id: str) -> None:
    """Remove all Qdrant points belonging to a specific content_id from the course collection."""
    col = collection_name(institute_id, course_id)
    try:
        _qdrant_request(
            "POST",
            f"/collections/{col}/points/delete",
            json={"filter": {"must": [{"key": "content_id", "match": {"value": content_id}}]}},
        )
        logger.info(f"Deleted Qdrant points for content {content_id!r} from collection {col!r}")
    except Exception as e:
        logger.error(f"Failed to delete content {content_id!r} from collection {col!r}: {e}", exc_info=True)


def collection_exists(col_name: str) -> bool:
    """Return True if the Qdrant collection exists, False otherwise."""
    try:
        _qdrant_request("GET", f"/collections/{col_name}")
        return True
    except Exception:
        return False


def search_course(institute_id: str, course_id: str, query: str, limit: int = 3) -> list[dict]:
    """Semantic search within a course's dedicated Qdrant collection.

    Returns an empty list (instead of raising) when the collection does not
    exist yet — this happens when no course material has been uploaded/indexed.

    Args:
        institute_id: Institute UUID (part of the collection name).
        course_id:    Course UUID (part of the collection name).
        query:        Natural-language search query.
        limit:        Maximum number of results to return.

    Returns:
        List of dicts with keys: content, page, title, score.
    """
    col = collection_name(institute_id, course_id)

    cache_key = (col, query, limit)
    cached = _cache_get(cache_key)
    if cached is not None:
        logger.debug(f"search_course cache hit for query={query!r} col={col!r}")
        return cached

    if not collection_exists(col):
        logger.info(f"Collection {col!r} does not exist — no KB indexed yet, returning empty results.")
        return []

    model = _get_embedding_model()
    query_vector = list(model.embed([query]))[0].tolist()

    result = _qdrant_request(
        "POST",
        f"/collections/{col}/points/search",
        json={
            "vector": query_vector,
            "limit": limit,
            "with_payload": True,
            "score_threshold": 0.35,
            "params": {"hnsw_ef": 64, "exact": False},
        },
    )
    results = [
        {
            "content": h["payload"].get("content", "")[:MAX_CONTENT_CHARS],
            "page": h["payload"].get("page"),
            "title": h["payload"].get("title", ""),
            "score": round(h["score"], 3),
        }
        for h in result.get("result", [])
    ]
    _cache_set(cache_key, results)
    return results
