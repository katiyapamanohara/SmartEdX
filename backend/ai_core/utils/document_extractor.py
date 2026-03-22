"""Extract plain text from PDF, DOCX and PPTX files."""

import io
from pathlib import Path


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return _from_pdf(file_bytes)
    if ext in (".docx", ".doc"):
        return _from_docx(file_bytes)
    if ext in (".pptx", ".ppt"):
        return _from_pptx(file_bytes)

    raise ValueError(f"Unsupported file type: {ext}")


def _from_pdf(data: bytes) -> str:
    import pdfplumber

    text_parts: list[str] = []
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                text_parts.append(t)
    return "\n\n".join(text_parts)


def _from_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _from_pptx(data: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        parts.append(line)
    return "\n".join(parts)
